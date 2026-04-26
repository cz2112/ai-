import base64
import json
import mimetypes
import subprocess
import tempfile
import wave
from pathlib import Path

import httpx

from app.core.config import settings


def _get_ai_credentials() -> tuple[str, str, str]:
    api_key = settings.DEEPSEEK_API_KEY or settings.AI_API_KEY
    base_url = (settings.AI_BASE_URL or "https://api.deepseek.com").rstrip("/")
    model = settings.LLM_MODEL or "deepseek-v4-flash"
    if not api_key:
        raise RuntimeError("DeepSeek API key is not configured")
    return api_key, base_url, model


def _get_glm_credentials() -> tuple[str, str]:
    api_key = settings.ZHIPU_API_KEY
    base_url = (settings.GLM_BASE_URL or "https://open.bigmodel.cn/api/paas/v4").rstrip("/")
    if not api_key:
        raise RuntimeError("ZHIPU_API_KEY is not configured")
    return api_key, base_url


def _raise_http_error(response: httpx.Response, label: str) -> None:
    try:
        response.raise_for_status()
    except httpx.HTTPStatusError as exc:
        detail = exc.response.text.strip().replace("\n", " ")
        raise RuntimeError(f"{label} failed with HTTP {exc.response.status_code}: {detail[:400]}") from exc


def _call_chat_completion(messages: list[dict], temperature: float, max_tokens: int) -> str:
    api_key, base_url, model = _get_ai_credentials()
    response = httpx.post(
        f"{base_url}/chat/completions",
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        json={
            "model": model,
            "messages": messages,
            "temperature": temperature,
            "max_tokens": max_tokens,
        },
        timeout=180,
    )
    _raise_http_error(response, "DeepSeek chat completion")

    payload = response.json()
    try:
        return payload["choices"][0]["message"]["content"]
    except (KeyError, IndexError, TypeError) as exc:
        raise RuntimeError("Unexpected DeepSeek response format") from exc


def _call_glm_chat_completion(messages: list[dict], temperature: float, max_tokens: int, model: str | None = None) -> str:
    api_key, base_url = _get_glm_credentials()
    response = httpx.post(
        f"{base_url}/chat/completions",
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        json={
            "model": model or settings.GLM_VISION_MODEL,
            "messages": messages,
            "temperature": temperature,
            "max_tokens": max_tokens,
        },
        timeout=180,
    )
    _raise_http_error(response, "GLM multimodal completion")

    payload = response.json()
    try:
        return payload["choices"][0]["message"]["content"]
    except (KeyError, IndexError, TypeError) as exc:
        raise RuntimeError("Unexpected GLM response format") from exc


def _encode_file_base64(file_path: str | Path) -> str:
    return base64.b64encode(Path(file_path).read_bytes()).decode("ascii")


def _mime_type_for_file(file_path: str | Path) -> str:
    mime_type, _ = mimetypes.guess_type(str(file_path))
    return mime_type or "application/octet-stream"


def _file_data_url(file_path: str | Path) -> str:
    return f"data:{_mime_type_for_file(file_path)};base64,{_encode_file_base64(file_path)}"


def _extract_pdf_text_locally(file_path: str) -> str:
    from PyPDF2 import PdfReader

    reader = PdfReader(file_path)
    text_parts = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
    return "\n".join(text_parts)


def _glm_layout_parsing(file_path: str, start_page_id: int | None = None, end_page_id: int | None = None) -> str:
    api_key, base_url = _get_glm_credentials()
    payload = {
        "model": settings.GLM_OCR_MODEL,
        "file": _file_data_url(file_path),
    }
    if start_page_id is not None:
        payload["start_page_id"] = start_page_id
    if end_page_id is not None:
        payload["end_page_id"] = end_page_id

    response = httpx.post(
        f"{base_url}/layout_parsing",
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
        },
        json=payload,
        timeout=300,
    )
    _raise_http_error(response, "GLM OCR")

    payload = response.json()
    markdown = payload.get("md_results", "")
    if isinstance(markdown, list):
        return "\n\n".join(part.strip() for part in markdown if isinstance(part, str) and part.strip())
    if isinstance(markdown, str):
        return markdown.strip()
    raise RuntimeError("Unexpected GLM OCR response format")


def extract_pdf_text(file_path: str) -> str:
    if not settings.ZHIPU_API_KEY:
        return _extract_pdf_text_locally(file_path)

    try:
        from PyPDF2 import PdfReader

        page_count = len(PdfReader(file_path).pages)
        if page_count <= 0:
            return ""

        sections = []
        for start_page in range(1, page_count + 1, 100):
            end_page = min(start_page + 99, page_count)
            content = _glm_layout_parsing(file_path, start_page_id=start_page, end_page_id=end_page)
            if content:
                sections.append(content)

        text = "\n\n".join(sections).strip()
        return text or _extract_pdf_text_locally(file_path)
    except Exception:
        return _extract_pdf_text_locally(file_path)


def extract_pptx_text(file_path: str) -> str:
    from pptx import Presentation

    presentation = Presentation(file_path)
    text_parts = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        slide_text = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_text.append(text)
        if slide_text:
            text_parts.append(f"[Slide {slide_number}]\n" + "\n".join(slide_text))
    return "\n\n".join(text_parts)


def extract_docx_text(file_path: str) -> str:
    from docx import Document

    document = Document(file_path)
    text_parts = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            text_parts.append(text)
    for table in document.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                text_parts.append(row_text)
    return "\n".join(text_parts)


def _extract_image_text_locally(file_path: str) -> str:
    try:
        import pytesseract
        from PIL import Image

        image = Image.open(file_path)
        return pytesseract.image_to_string(image).strip()
    except Exception:
        return "[OCR processing failed or pytesseract is unavailable]"


def extract_image_text(file_path: str) -> str:
    if not settings.ZHIPU_API_KEY:
        return _extract_image_text_locally(file_path)

    try:
        text = _glm_layout_parsing(file_path)
        return text or _extract_image_text_locally(file_path)
    except Exception:
        return _extract_image_text_locally(file_path)


def _get_ffmpeg_executable() -> str:
    from imageio_ffmpeg import get_ffmpeg_exe

    return get_ffmpeg_exe()


def _run_ffmpeg(arguments: list[str], label: str) -> None:
    ffmpeg_path = _get_ffmpeg_executable()
    process = subprocess.run(
        [ffmpeg_path, "-hide_banner", "-loglevel", "error", *arguments],
        capture_output=True,
        text=True,
        check=False,
    )
    if process.returncode != 0:
        stderr = (process.stderr or process.stdout or "").strip().replace("\n", " ")
        raise RuntimeError(f"{label} failed: {stderr[:400]}")


def _convert_media_to_wav(input_path: str, output_path: str) -> None:
    _run_ffmpeg(
        ["-y", "-i", input_path, "-vn", "-ac", "1", "-ar", "16000", output_path],
        "Audio conversion",
    )


def _split_wav_file(wav_path: str, chunk_seconds: int) -> list[str]:
    chunk_paths = []
    with wave.open(wav_path, "rb") as source:
        params = source.getparams()
        frame_rate = source.getframerate()
        total_frames = source.getnframes()
        frames_per_chunk = max(int(frame_rate * chunk_seconds), frame_rate)

        chunk_index = 0
        while source.tell() < total_frames:
            chunk_index += 1
            frames = source.readframes(frames_per_chunk)
            if not frames:
                break

            chunk_file = tempfile.NamedTemporaryFile(delete=False, suffix=f"-chunk-{chunk_index}.wav")
            chunk_file.close()
            with wave.open(chunk_file.name, "wb") as target:
                target.setparams(params)
                target.writeframes(frames)
            chunk_paths.append(chunk_file.name)

    return chunk_paths


def _transcribe_audio_chunk(chunk_path: str, prompt: str = "") -> str:
    api_key, base_url = _get_glm_credentials()
    data = {
        "model": settings.GLM_ASR_MODEL,
        "stream": "false",
    }
    if prompt.strip():
        data["prompt"] = prompt[-2000:]

    with open(chunk_path, "rb") as audio_file:
        response = httpx.post(
            f"{base_url}/audio/transcriptions",
            headers={"Authorization": f"Bearer {api_key}"},
            data=data,
            files={"file": (Path(chunk_path).name, audio_file, "audio/wav")},
            timeout=300,
        )
    _raise_http_error(response, "GLM audio transcription")

    payload = response.json()
    text = payload.get("text", "")
    if not isinstance(text, str):
        raise RuntimeError("Unexpected GLM ASR response format")
    return text.strip()


def _transcribe_audio_with_glm(file_path: str) -> str:
    converted_wav = tempfile.NamedTemporaryFile(delete=False, suffix=".wav")
    converted_wav.close()
    chunk_paths: list[str] = []

    try:
        _convert_media_to_wav(file_path, converted_wav.name)
        chunk_paths = _split_wav_file(converted_wav.name, settings.GLM_ASR_CHUNK_SECONDS)
        transcripts = []
        running_context = ""

        for chunk_path in chunk_paths:
            transcript = _transcribe_audio_chunk(chunk_path, prompt=running_context)
            if transcript:
                transcripts.append(transcript)
                running_context = "\n".join(transcripts)[-4000:]

        final_text = "\n".join(part for part in transcripts if part).strip()
        if not final_text:
            raise RuntimeError("GLM ASR returned empty transcription")
        return final_text
    finally:
        for chunk_path in chunk_paths:
            Path(chunk_path).unlink(missing_ok=True)
        Path(converted_wav.name).unlink(missing_ok=True)


def transcribe_audio(file_path: str, language: str = "en") -> str:
    del language
    return _transcribe_audio_with_glm(file_path)


def _extract_video_audio_wav(video_path: str, output_path: str) -> None:
    _run_ffmpeg(
        ["-y", "-i", video_path, "-vn", "-ac", "1", "-ar", "16000", output_path],
        "Video audio extraction",
    )


def _extract_video_frames(video_path: str, frame_count: int) -> list[Path]:
    from imageio_ffmpeg import count_frames_and_secs

    _, duration_seconds = count_frames_and_secs(video_path)
    duration_seconds = max(float(duration_seconds), 1.0)
    frame_paths = []

    with tempfile.TemporaryDirectory() as temp_dir:
        temp_dir_path = Path(temp_dir)
        timestamps = [duration_seconds * index / (frame_count + 1) for index in range(1, frame_count + 1)]
        if not timestamps:
            timestamps = [0.0]

        for index, timestamp in enumerate(timestamps, start=1):
            frame_path = temp_dir_path / f"frame-{index}.jpg"
            try:
                _run_ffmpeg(
                    ["-y", "-ss", f"{timestamp:.3f}", "-i", video_path, "-frames:v", "1", "-q:v", "2", str(frame_path)],
                    "Video frame extraction",
                )
            except RuntimeError:
                continue

            if frame_path.exists() and frame_path.stat().st_size > 0:
                copied = Path(tempfile.NamedTemporaryFile(delete=False, suffix=f"-frame-{index}.jpg").name)
                copied.write_bytes(frame_path.read_bytes())
                frame_paths.append(copied)

    return frame_paths


def _describe_video_frames(frame_paths: list[Path], filename: str) -> str:
    content = [
        {
            "type": "text",
            "text": (
                f"These images are sampled frames from the video file {filename}. "
                "Describe the visible content in time order. Extract readable on-screen text, slide titles, equations, "
                "labels, UI text, tables, and scene changes. Keep the output factual and compact."
            ),
        }
    ]
    for frame_path in frame_paths:
        content.append(
            {
                "type": "image_url",
                "image_url": {"url": _file_data_url(frame_path)},
            }
        )

    return _call_glm_chat_completion(
        [{"role": "user", "content": content}],
        temperature=0.1,
        max_tokens=2048,
        model=settings.GLM_VISION_MODEL,
    ).strip()


def extract_video_text(file_path: str) -> str:
    if not settings.ZHIPU_API_KEY:
        raise RuntimeError("Video analysis requires ZHIPU_API_KEY")

    audio_text = ""
    visual_text = ""
    audio_path = tempfile.NamedTemporaryFile(delete=False, suffix=".wav")
    audio_path.close()
    frame_paths: list[Path] = []

    try:
        try:
            _extract_video_audio_wav(file_path, audio_path.name)
            audio_text = _transcribe_audio_with_glm(audio_path.name)
        except Exception:
            audio_text = ""

        frame_paths = _extract_video_frames(file_path, settings.GLM_VIDEO_FRAME_COUNT)
        if frame_paths:
            visual_text = _describe_video_frames(frame_paths, Path(file_path).name)

        sections = []
        if audio_text:
            sections.append("Audio transcript:\n" + audio_text)
        if visual_text:
            sections.append("Visual analysis:\n" + visual_text)

        if not sections:
            raise RuntimeError("Unable to extract either audio transcript or visual notes from the video")

        return "\n\n".join(sections).strip()
    finally:
        Path(audio_path.name).unlink(missing_ok=True)
        for frame_path in frame_paths:
            frame_path.unlink(missing_ok=True)


def detect_language(text: str) -> str:
    sample = text[:2000]
    prompt = (
        "Detect the primary language of the following text. "
        "Return only the ISO 639-1 language code, for example en, zh, ja, ko, fr, de, es."
    )
    code = _chat(prompt, sample).strip().lower()[:2]
    return code if len(code) == 2 else "en"


def _strip_json_fences(raw: str) -> str:
    cleaned = raw.strip()
    if cleaned.startswith("```"):
        cleaned = cleaned.split("\n", 1)[1] if "\n" in cleaned else cleaned[3:]
    if cleaned.endswith("```"):
        cleaned = cleaned[:-3]
    return cleaned.strip()


def _load_json(raw: str):
    cleaned = _strip_json_fences(raw)
    try:
        return json.loads(cleaned)
    except json.JSONDecodeError:
        start_positions = [index for index in (cleaned.find("{"), cleaned.find("[")) if index != -1]
        if not start_positions:
            raise
        start = min(start_positions)
        end = max(cleaned.rfind("}"), cleaned.rfind("]"))
        if end == -1 or end <= start:
            raise
        return json.loads(cleaned[start : end + 1])


def _chat(prompt: str, text: str, temperature: float = 0.3, max_tokens: int = 4096) -> str:
    max_chars = 24000
    if len(text) > max_chars:
        text = text[:max_chars] + "\n...[truncated]"

    return _call_chat_completion(
        [
            {"role": "system", "content": prompt},
            {"role": "user", "content": text},
        ],
        temperature=temperature,
        max_tokens=max_tokens,
    )


def _lang_instruction(lang: str) -> str:
    if lang == "zh":
        return "Respond in Chinese. "
    if lang == "ja":
        return "Respond in Japanese. "
    if lang == "ko":
        return "Respond in Korean. "
    if lang != "en":
        return f"Respond in the same language as the input material (language code: {lang}). "
    return ""


def generate_summary(text: str, lang: str = "en") -> str:
    prompt = (
        f"{_lang_instruction(lang)}"
        "You are a study assistant. Given the following study material, "
        "write a concise summary of 200 to 400 words that captures the main ideas, "
        "arguments, and conclusions. Use clear academic language."
    )
    return _chat(prompt, text)


def generate_key_concepts(text: str, lang: str = "en") -> list[dict]:
    prompt = (
        f"{_lang_instruction(lang)}"
        "You are a study assistant. Extract 5 to 10 key concepts from the following material. "
        "Return only a JSON array. Each element must contain title, description, and citation fields. "
        "The citation must be a short supporting excerpt from the original text."
    )
    return _load_json(_chat(prompt, text))


def generate_flashcards(text: str, lang: str = "en") -> list[dict]:
    prompt = (
        f"{_lang_instruction(lang)}"
        "You are a study assistant. Generate 8 to 15 flashcards from the following material. "
        "Return only a JSON array with question and answer fields."
    )
    return _load_json(_chat(prompt, text))


def chat_with_context(messages: list[dict], context_text: str, lang: str = "en") -> str:
    system_prompt = (
        f"{_lang_instruction(lang)}"
        "You are a study assistant. Answer questions based on the following study material. "
        "If the answer is not in the material, say so clearly.\n\n"
        f"Study Material:\n{context_text[:20000]}"
    )
    return _call_chat_completion(
        [{"role": "system", "content": system_prompt}] + messages,
        temperature=0.4,
        max_tokens=2048,
    )


def generate_knowledge_graph(text: str, lang: str = "en") -> dict:
    prompt = (
        f"{_lang_instruction(lang)}"
        "Extract a knowledge graph from the following study material. "
        "Return only a JSON object with nodes and edges arrays. "
        "Each node must contain id, label, and group. "
        "Each edge must contain source, target, and label. "
        "Limit the graph to 15 to 25 nodes."
    )
    return _load_json(_chat(prompt, text))


def generate_learning_path(summaries: list[str], known_concepts: list[str], lang: str = "en") -> list[dict]:
    prompt = (
        f"{_lang_instruction(lang)}"
        "Based on the student's study materials and known concepts, "
        "suggest a learning path with 5 to 8 steps. "
        "Return only a JSON array. Each item should contain step, title, description, and resources."
    )
    combined = "Materials:\n" + "\n---\n".join(summaries[:5])
    combined += "\n\nKnown concepts:\n" + ", ".join(known_concepts[:30])
    return _load_json(_chat(prompt, combined))
